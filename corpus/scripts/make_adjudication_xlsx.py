"""Render the adjudication CSV as a formatted .xlsx for easy manual editing.

Nice-to-edit features: frozen header + string column, auto-filter, sensible
column widths with wrapped affiliation strings, the `correct_ror` column
highlighted, disagreement rows tinted so they stand out from the pre-filled
agreement rows, and a short instructions row at the top.

Usage:
    python scripts/make_adjudication_xlsx.py
"""
from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parent.parent

WIDTHS = {
    "string": 52, "n_papers": 9, "agree": 8,
    "multi_ror": 13, "multi_name": 34, "multi_country": 9, "multi_score": 9,
    "single_ror": 13, "single_name": 34, "single_country": 9, "single_score": 9,
    "correct_ror": 18, "notes": 26,
}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--csv", type=Path,
                    default=REPO / "reports/relink_compare/adjudication_sheet.csv")
    ap.add_argument("--out", type=Path,
                    default=REPO / "reports/relink_compare/adjudication_sheet.xlsx")
    args = ap.parse_args()

    with args.csv.open(newline="", encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    header, data = rows[0], rows[1:]
    ci = {h: i for i, h in enumerate(header)}

    wb = Workbook()
    ws = wb.active
    ws.title = "adjudication"

    # Instructions banner (row 1), header (row 2), data from row 3.
    instr = ("Fill correct_ror for the DISAGREEMENT rows (top, yellow). Copy multi_ror "
             "or single_ror, or type another bare ROR id; ';'-separate several; 'none' "
             "if no ROR applies. AGREE rows (bottom) are pre-filled — spot-check only.")
    ws.cell(row=1, column=1, value=instr)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(header))
    ws.cell(row=1, column=1).font = Font(italic=True, color="555555")
    ws.cell(row=1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 42

    head_fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(bold=True, color="FFFFFF")
    for j, h in enumerate(header, start=1):
        c = ws.cell(row=2, column=j, value=h)
        c.fill = head_fill; c.font = head_font
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(h, 12)

    dis_fill = PatternFill("solid", fgColor="FFF3CD")      # yellow-ish: do these
    agree_fill = PatternFill("solid", fgColor="EAF3EA")    # green-ish: pre-filled
    correct_dis = PatternFill("solid", fgColor="FFE08A")   # stronger yellow for input cell
    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for r, row in enumerate(data, start=3):
        is_dis = (row[ci["agree"]].strip() == "")  # disagreements have blank 'agree'
        fill = dis_fill if is_dis else agree_fill
        for j, val in enumerate(row, start=1):
            c = ws.cell(row=r, column=j, value=val)
            c.fill = fill
            c.border = border
            c.alignment = Alignment(wrap_text=(header[j-1] in ("string", "multi_name",
                                                               "single_name", "notes")),
                                    vertical="top")
        # highlight the editable cell on disagreement rows
        cc = ci["correct_ror"] + 1
        if is_dis:
            ws.cell(row=r, column=cc).fill = correct_dis

    ws.freeze_panes = "B3"  # keep header rows + the 'string' column visible
    ws.auto_filter.ref = f"A2:{get_column_letter(len(header))}{len(data)+2}"

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    n_dis = sum(1 for row in data if row[ci["agree"]].strip() == "")
    print(f"# wrote {args.out}  ({len(data)} rows, {n_dis} disagreements to fill)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
