"""Score CSO topic-labeling precision from the annotated validation sheet.

Reads applies? (yes/no/unsure) + source (in-text/semantic) from
cso_validation.xlsx and reports precision overall and by source, under two
conventions for 'unsure' (excluded, and counted-as-no), with Wilson 95% CIs.
"""
from __future__ import annotations

import argparse
import math
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parent.parent


def wilson(k, n, z=1.96):
    if n == 0:
        return (float("nan"), float("nan"))
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return (c - h, c + h)


def line(label, yes, no, uns):
    dec = yes + no
    tot = yes + no + uns
    p_excl = yes / dec if dec else float("nan")
    lo1, hi1 = wilson(yes, dec)
    p_incl = yes / tot if tot else float("nan")     # unsure = no
    lo2, hi2 = wilson(yes, tot)
    return (f"| {label} | {yes} | {no} | {uns} | {p_excl:.3f} [{lo1:.3f}, {hi1:.3f}] "
            f"| {p_incl:.3f} [{lo2:.3f}, {hi2:.3f}] |")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--sheet", type=Path,
                    default=REPO / "reports/cso_validation/cso_validation.xlsx")
    ap.add_argument("--out", type=Path,
                    default=REPO / "reports/cso_validation/cso_precision.md")
    args = ap.parse_args()

    wb = load_workbook(args.sheet)
    ws = wb.active
    hdr = [c.value for c in ws[2]]
    I = {h: i for i, h in enumerate(hdr)}
    c = Counter()
    for r in range(3, ws.max_row + 1):
        a = ws.cell(r, I["applies?"] + 1).value
        src = ws.cell(r, I["source"] + 1).value
        if a in (None, ""):
            continue
        v = str(a).strip().lower()
        if v not in ("yes", "no", "unsure"):
            continue
        c[("all", v)] += 1
        c[(src, v)] += 1

    md = []
    md.append("# CSO topic-labeling precision (annotated sample)\n")
    md.append("| Slice | yes | no | unsure | precision (excl. unsure) | precision (unsure=no) |")
    md.append("|---|---:|---:|---:|---|---|")
    for label in ("all", "in-text", "semantic"):
        md.append(line(label, c[(label, "yes")], c[(label, "no")], c[(label, "unsure")]))
    md.append("")
    tot = c[("all", "yes")] + c[("all", "no")] + c[("all", "unsure")]
    uns = c[("all", "unsure")]
    md.append(f"- unsure rate (subjectivity): {uns}/{tot} = {uns/tot:.1%}")
    md.append("- Wilson 95% CIs; single annotator, small sample — treat as indicative.")
    args.out.write_text("\n".join(md))
    print("\n".join(md))
    print(f"\n# wrote {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
