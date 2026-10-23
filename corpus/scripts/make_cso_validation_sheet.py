"""Build a CSO topic-precision annotation sheet (read-once, tick-down).

For a sample of papers, lists each paper's `topics.cso.union` labels (the
operative fine-grained set) with title + authors + abstract shown once per
paper, a `source` tag (in-text vs semantic-only), and a dropdown `applies?`
column. You judge only the union topics; the hierarchical `enhanced` parents
are validated automatically later (an enhanced topic is correct iff it is a
CSO-ancestor of a union topic you marked correct — needs the CSO ontology at
scoring time, not here).

Sampling: cluster by paper (sample N papers, annotate all their union topics) so
you read each abstract once. By default restricts to papers in the affiliation
gold set (--gold-csv) and pulls authors from there; --from-enriched-pool samples
any enriched paper (use when the gold papers' enriched JSONs aren't on this host).

Outputs a standalone xlsx + csv; --append-to adds a 'cso_validation' tab to an
existing workbook (e.g. the ROR adjudication sheet) so it's one file.

Usage (real, on the host with the full enriched corpus):
    python scripts/make_cso_validation_sheet.py \
        --enriched-dir /share/.../ror_parsed_affiliations_90_with_topics_and_bibmeta \
        --n-papers 15
"""
from __future__ import annotations

import argparse
import csv
import json
import random
from pathlib import Path
from typing import Dict, List, Optional

REPO = Path(__file__).resolve().parent.parent
DEF_GOLD = Path("/home/moritz/Downloads/11688_Singing_from_the_Same_AN_Data/"
                "sample_data/affiliation_ground_truth.csv")
DEF_ENR = Path("/home/moritz/Downloads/11688_Singing_from_the_Same_AN_Data/"
               "sample_data/ror_parsed_affiliations_90_with_topics_and_bibmeta")


def gold_authors(gold_csv: Path) -> Dict[str, str]:
    out: Dict[str, str] = {}
    if not gold_csv.exists():
        return out
    for r in csv.DictReader(gold_csv.open(newline="", encoding="utf-8")):
        pid = (r.get("anthology_id") or "").strip()
        try:
            authors = json.loads(r.get("authors_raw_json") or "[]")
        except json.JSONDecodeError:
            authors = []
        names = ", ".join(a.get("name", "") for a in authors if a.get("name"))
        if pid:
            out[pid] = names
    return out


def collect(enriched_dir: Path, restrict: Optional[set], authors_map: Dict[str, str],
            from_pool: bool):
    papers = []
    for fp in sorted(enriched_dir.glob("*.json")):
        pid = fp.stem
        if restrict is not None and not from_pool and pid not in restrict:
            continue
        try:
            d = json.loads(fp.read_text(encoding="utf-8", errors="replace"))
        except json.JSONDecodeError:
            continue
        cso = (d.get("topics") or {}).get("cso") or {}
        union = cso.get("union") or []
        abstract = d.get("abstract") or ""
        if not union or not abstract:
            continue
        syntactic = {str(x).lower() for x in (cso.get("syntactic") or [])}
        authors = authors_map.get(pid) or ", ".join(
            a.get("name", "") for a in (d.get("authors") or []) if a.get("name"))
        papers.append({
            "paper_id": pid,
            "title": d.get("title") or "",
            "authors": authors,
            "abstract": abstract,
            "topics": [(t, "in-text" if str(t).lower() in syntactic else "semantic")
                       for t in union],
        })
    return papers


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--gold-csv", type=Path, default=DEF_GOLD)
    ap.add_argument("--enriched-dir", type=Path, default=DEF_ENR)
    ap.add_argument("--n-papers", type=int, default=15)
    ap.add_argument("--target-pairs", type=int, default=100,
                    help="Stop adding papers once this many (paper,topic) pairs reached.")
    ap.add_argument("--seed", type=int, default=42)
    ap.add_argument("--from-enriched-pool", action="store_true",
                    help="Sample any enriched paper (ignore gold restriction).")
    ap.add_argument("--append-to", type=Path, default=None,
                    help="Existing .xlsx to add a 'cso_validation' tab to.")
    ap.add_argument("--out-xlsx", type=Path,
                    default=REPO / "reports/cso_validation/cso_validation.xlsx")
    ap.add_argument("--out-csv", type=Path,
                    default=REPO / "reports/cso_validation/cso_validation.csv")
    args = ap.parse_args()

    authors_map = gold_authors(args.gold_csv)
    restrict = set(authors_map.keys()) if authors_map else None
    papers = collect(args.enriched_dir, restrict, authors_map, args.from_enriched_pool)
    print(f"# candidate papers with union+abstract"
          f"{' (gold-restricted)' if (restrict and not args.from_enriched_pool) else ''}: "
          f"{len(papers)}", flush=True)
    if not papers:
        print("# none available on this host — run where the enriched corpus lives, "
              "or pass --from-enriched-pool", flush=True)
        return 1

    rnd = random.Random(args.seed)
    rnd.shuffle(papers)
    chosen, pairs = [], 0
    for p in papers:
        chosen.append(p)
        pairs += len(p["topics"])
        if len(chosen) >= args.n_papers or pairs >= args.target_pairs:
            break
    print(f"# sampled {len(chosen)} papers / {pairs} (paper,topic) pairs "
          f"(seed {args.seed})", flush=True)

    cols = ["paper_id", "title", "authors", "abstract", "cso_topic", "source",
            "applies?", "notes"]
    rows = []
    for p in chosen:
        for i, (topic, src) in enumerate(p["topics"]):
            rows.append([
                p["paper_id"],
                p["title"] if i == 0 else "",
                p["authors"] if i == 0 else "",
                p["abstract"] if i == 0 else "",
                topic, src, "", "",
            ])

    args.out_csv.parent.mkdir(parents=True, exist_ok=True)
    with args.out_csv.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(cols)
        w.writerows(rows)

    print(f"# wrote {args.out_csv}")
    try:
        _write_xlsx(cols, rows, chosen, args.out_xlsx, args.append_to)
        print(f"# wrote {args.out_xlsx if not args.append_to else args.append_to} "
              f"(tab 'cso_validation')")
    except ImportError:
        print("# openpyxl not installed — CSV written, xlsx skipped "
              "(pip install openpyxl to get the formatted sheet)")
    return 0


def _write_xlsx(cols, rows, chosen, out_xlsx: Path, append_to: Optional[Path]):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    if append_to and append_to.exists():
        wb = load_workbook(append_to)
        if "cso_validation" in wb.sheetnames:
            del wb["cso_validation"]
        ws = wb.create_sheet("cso_validation")
        target = append_to
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "cso_validation"
        target = out_xlsx

    widths = {"paper_id": 20, "title": 40, "authors": 30, "abstract": 70,
              "cso_topic": 30, "source": 11, "applies?": 11, "notes": 22}
    head_fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(bold=True, color="FFFFFF")
    ws.cell(row=1, column=1, value=("Read each paper's abstract once, then set applies? "
            "(yes/no/unsure) for each cso_topic. In-text topics are usually yes; scrutinise "
            "'semantic' ones. Parent/hierarchy topics are validated automatically later."))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.cell(row=1, column=1).font = Font(italic=True, color="555555")
    ws.cell(row=1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 40
    for j, h in enumerate(cols, 1):
        c = ws.cell(row=2, column=j, value=h)
        c.fill = head_fill; c.font = head_font
        ws.column_dimensions[get_column_letter(j)].width = widths.get(h, 12)

    tintA = PatternFill("solid", fgColor="F2F6FB")
    tintB = PatternFill("solid", fgColor="FFFFFF")
    input_fill = PatternFill("solid", fgColor="FFF3CD")
    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap_cols = {"title", "authors", "abstract", "cso_topic", "notes"}

    pid_order = []
    for r in rows:
        if r[0] not in pid_order:
            pid_order.append(r[0])
    block = {pid: (i % 2 == 0) for i, pid in enumerate(pid_order)}

    ap_col = cols.index("applies?") + 1
    for ri, row in enumerate(rows, start=3):
        fill = tintA if block[row[0]] else tintB
        for j, val in enumerate(row, 1):
            c = ws.cell(row=ri, column=j, value=val)
            c.fill = fill
            c.border = border
            c.alignment = Alignment(wrap_text=(cols[j-1] in wrap_cols), vertical="top")
        ws.cell(row=ri, column=ap_col).fill = input_fill

    dv = DataValidation(type="list", formula1='"yes,no,unsure"', allow_blank=True)
    ws.add_data_validation(dv)
    col_letter = get_column_letter(ap_col)
    dv.add(f"{col_letter}3:{col_letter}{len(rows)+2}")

    ws.freeze_panes = "E3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{len(rows)+2}"
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)


if __name__ == "__main__":
    raise SystemExit(main())
